"""Document parser — extracts text blocks with heading detection.

Supports: PDF (PyMuPDF), DOCX, PPTX, TXT/MD.
Falls back to Tesseract OCR for scanned/image-heavy PDF pages.
"""

import io
import logging
import statistics
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF

logger = logging.getLogger(__name__)


@dataclass
class Block:
    """A single text block from a document page."""
    page: int
    text: str
    is_heading: bool = False


@dataclass
class ParsedDoc:
    """Result of parsing a document into blocks."""
    blocks: list[Block] = field(default_factory=list)
    page_count: int = 0


# ---------------------------------------------------------------------------
# OCR fallback
# ---------------------------------------------------------------------------

def _ocr_page(page: fitz.Page) -> str:
    """Rasterise a page and OCR it with Tesseract."""
    try:
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img).strip()
    except Exception as exc:
        logger.warning("OCR failed for page %d: %s", page.number + 1, exc)
        return ""


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _collect_font_sizes(doc: fitz.Document) -> list[float]:
    """Collect all non-empty span font sizes across the document."""
    sizes: list[float] = []
    for page in doc:
        for block in page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]:
            if block.get("type") != 0:  # skip image blocks
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span["text"].strip():
                        sizes.append(span["size"])
    return sizes


def _clean_text(text: str) -> str:
    """Normalise whitespace: collapse runs, strip."""
    import re
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(path: str) -> ParsedDoc:
    """Parse a PDF into blocks, detecting headings by font size."""
    with fitz.open(path) as doc:
        parsed = ParsedDoc(page_count=len(doc))

        # Determine the median body font size → heading threshold
        sizes = _collect_font_sizes(doc)
        if not sizes:
            # Entirely scanned document — OCR every page
            for pno, page in enumerate(doc, start=1):
                ocr_text = _ocr_page(page)
                if ocr_text:
                    parsed.blocks.append(Block(pno, _clean_text(ocr_text)))
            return parsed

        body_size = statistics.median(sizes)
        heading_threshold = body_size * 1.15

        for pno, page in enumerate(doc, start=1):
            page_text = page.get_text("text").strip()

            # If barely any extractable text, try OCR
            if len(page_text) < 30:
                ocr_text = _ocr_page(page)
                if ocr_text:
                    parsed.blocks.append(Block(pno, _clean_text(ocr_text)))
                continue

            # Extract structured blocks
            for block in page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]:
                if block.get("type") != 0:
                    continue

                lines_text: list[str] = []
                max_font_size = 0.0

                for line in block.get("lines", []):
                    line_str = "".join(span["text"] for span in line.get("spans", [])).strip()
                    if line_str:
                        lines_text.append(line_str)
                        for span in line.get("spans", []):
                            if span["text"].strip():
                                max_font_size = max(max_font_size, span["size"])

                text = " ".join(lines_text).strip()
                if not text:
                    continue

                # A heading is large font + short text
                is_heading = (
                    max_font_size >= heading_threshold
                    and len(text.split()) <= 15
                )
                parsed.blocks.append(Block(pno, _clean_text(text), is_heading))

        return parsed


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def parse_docx(path: str) -> ParsedDoc:
    """Parse a DOCX file. Heading styles are detected from paragraph styles."""
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    parsed = ParsedDoc(page_count=1)  # DOCX doesn't have reliable page info

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        is_heading = para.style and para.style.name.lower().startswith("heading")
        parsed.blocks.append(Block(1, text, is_heading))

    return parsed


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def parse_pptx(path: str) -> ParsedDoc:
    """Parse a PPTX file. Each slide becomes a page; slide titles are headings."""
    from pptx import Presentation

    prs = Presentation(path)
    parsed = ParsedDoc(page_count=len(prs.slides))

    for slide_num, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            is_title = (slide.shapes.title is not None and shape == slide.shapes.title)
            parsed.blocks.append(Block(slide_num, text, is_title))

    return parsed


# ---------------------------------------------------------------------------
# TXT / MD
# ---------------------------------------------------------------------------

def parse_text(path: str) -> ParsedDoc:
    """Parse a plain text or Markdown file."""
    content = Path(path).read_text(encoding="utf-8", errors="ignore")
    parsed = ParsedDoc(page_count=1)

    for para in content.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        # Markdown headings
        is_heading = para.startswith("#")
        display_text = para.lstrip("#").strip() if is_heading else para
        if display_text:
            parsed.blocks.append(Block(1, display_text, is_heading))

    return parsed


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------

PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_file(path: str) -> ParsedDoc:
    """Parse any supported file type into a ParsedDoc."""
    ext = Path(path).suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: {ext}")
    logger.info("Parsing %s with %s parser", path, ext)
    return parser(path)
